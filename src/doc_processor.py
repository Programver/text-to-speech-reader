"""
Document extraction and processing module.
Supports PDF (PyMuPDF), DOCX (python-docx), and plain text (TXT).
"""

import io
from typing import Dict, Any, Tuple


def get_text_metrics(text: str) -> Dict[str, Any]:
    """Calculates character count, word count, and estimated narration duration."""
    clean_text = text.strip()
    char_count = len(clean_text)
    char_count_no_spaces = len("".join(clean_text.split()))
    words = clean_text.split()
    word_count = len(words)
    # Average speaking rate ~130-150 words per minute
    est_duration_minutes = word_count / 140.0 if word_count > 0 else 0.0

    return {
        "char_count": char_count,
        "char_count_no_spaces": char_count_no_spaces,
        "word_count": word_count,
        "est_duration_minutes": round(est_duration_minutes, 1),
        "est_duration_seconds": int(est_duration_minutes * 60),
    }


def extract_from_pdf(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
    """
    Extracts readable text from a PDF byte stream using PyMuPDF (fitz).
    Detects scanned/image-only PDFs.
    """
    import pymupdf as fitz

    doc = fitz.open(stream=file_bytes, filetype="pdf")
    num_pages = len(doc)
    extracted_pages = []

    for page_idx in range(num_pages):
        page = doc[page_idx]
        page_text = page.get_text("text")
        if page_text.strip():
            extracted_pages.append(page_text.strip())

    doc.close()
    full_text = "\n\n".join(extracted_pages).strip()

    metadata = {
        "file_type": "PDF",
        "num_pages": num_pages,
        "pages_with_text": len(extracted_pages),
        "is_scanned": len(full_text) == 0 and num_pages > 0,
    }

    return full_text, metadata


def extract_from_docx(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
    """Extracts readable text from a DOCX byte stream using python-docx."""
    import docx

    file_stream = io.BytesIO(file_bytes)
    doc = docx.Document(file_stream)

    paragraphs_text = []
    for p in doc.paragraphs:
        if p.text.strip():
            paragraphs_text.append(p.text.strip())

    # Also extract any text inside tables
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                paragraphs_text.append(" | ".join(row_text))

    full_text = "\n\n".join(paragraphs_text).strip()

    metadata = {
        "file_type": "DOCX",
        "paragraphs_count": len(doc.paragraphs),
        "tables_count": len(doc.tables),
        "is_scanned": False,
    }

    return full_text, metadata


def extract_from_txt(file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
    """Extracts text from plain text bytes trying common encodings."""
    encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252"]
    decoded_text = ""
    used_encoding = "utf-8"

    for enc in encodings:
        try:
            decoded_text = file_bytes.decode(enc)
            used_encoding = enc
            break
        except UnicodeDecodeError:
            continue

    clean_text = decoded_text.strip()
    metadata = {
        "file_type": "TXT",
        "encoding": used_encoding,
        "is_scanned": False,
    }

    return clean_text, metadata


def extract_document_text(file_bytes: bytes, file_name: str) -> Tuple[str, Dict[str, Any], str]:
    """
    Main entry point for document text extraction.
    Returns: (extracted_text, metadata, error_message)
    """
    lower_name = file_name.lower()
    try:
        if lower_name.endswith(".pdf"):
            text, meta = extract_from_pdf(file_bytes)
            if meta.get("is_scanned"):
                return (
                    "",
                    meta,
                    "This PDF appears to be scanned or image-only without an embedded text layer. Please use a text-based PDF or copy-paste text directly.",
                )
            return text, meta, ""
        elif lower_name.endswith(".docx"):
            text, meta = extract_from_docx(file_bytes)
            return text, meta, ""
        elif lower_name.endswith(".txt"):
            text, meta = extract_from_txt(file_bytes)
            return text, meta, ""
        else:
            return "", {}, f"Unsupported file format: {file_name}. Supported formats are .pdf, .docx, and .txt"
    except Exception as e:
        return "", {}, f"Error processing file '{file_name}': {str(e)}"
