"""
Unit test for document text extraction (PDF, DOCX, TXT).
Creates test sample documents dynamically and checks extraction accuracy.
"""

import os
from src.doc_processor import extract_document_text, get_text_metrics


def test_txt_extraction():
    os.makedirs("test_samples", exist_ok=True)
    txt_path = "test_samples/sample.txt"
    sample_content = "This is a plain text sample file for AudioReader testing.\nIt has multiple lines and paragraphs."
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(sample_content)

    with open(txt_path, "rb") as f:
        text, meta, err = extract_document_text(f.read(), "sample.txt")

    assert err == "", f"TXT extraction returned error: {err}"
    assert "plain text sample file" in text
    assert meta["file_type"] == "TXT"
    print("✓ TXT extraction test passed")


def test_docx_extraction():
    import docx
    os.makedirs("test_samples", exist_ok=True)
    docx_path = "test_samples/sample.docx"
    doc = docx.Document()
    doc.add_heading("AudioReader DOCX Sample", 0)
    doc.add_paragraph("First paragraph testing python-docx extraction.")
    doc.add_paragraph("Second paragraph with multilingual content: नमस्ते भारत।")
    doc.save(docx_path)

    with open(docx_path, "rb") as f:
        text, meta, err = extract_document_text(f.read(), "sample.docx")

    assert err == "", f"DOCX extraction returned error: {err}"
    assert "First paragraph testing" in text
    assert "नमस्ते भारत" in text
    assert meta["file_type"] == "DOCX"
    print("✓ DOCX extraction test passed")


def test_pdf_extraction():
    import pymupdf as fitz
    os.makedirs("test_samples", exist_ok=True)
    pdf_path = "test_samples/sample.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), "AudioReader PDF Sample Content.\nTesting PyMuPDF text layer extraction.")
    doc.save(pdf_path)
    doc.close()

    with open(pdf_path, "rb") as f:
        text, meta, err = extract_document_text(f.read(), "sample.pdf")

    assert err == "", f"PDF extraction returned error: {err}"
    assert "Testing PyMuPDF text layer" in text
    assert meta["file_type"] == "PDF"
    print("✓ PDF extraction test passed")


def test_metrics_calculation():
    text = "One two three four five six seven eight nine ten."
    metrics = get_text_metrics(text)
    assert metrics["word_count"] == 10
    assert metrics["char_count"] == len(text)
    print("✓ Metrics calculation test passed")


if __name__ == "__main__":
    test_txt_extraction()
    test_docx_extraction()
    test_pdf_extraction()
    test_metrics_calculation()
    print("All document processor tests passed!")
